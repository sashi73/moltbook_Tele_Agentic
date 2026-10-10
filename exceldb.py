import os
import sqlite3
from datetime import datetime

try:
    from openpyxl import Workbook
    from openpyxl.styles import Font, Alignment, PatternFill
    from openpyxl.utils import get_column_letter
except ImportError:
    print("❌ openpyxl is not installed.")
    print("   Install it with: pip install openpyxl")
    raise SystemExit(1)

DB_NAME = "agent_memory.db"
ANSWER_TRUNCATE = 500   # max chars per answer in the export
TITLE_TRUNCATE = 500    # max chars per Moltbook title/body


def truncate(text, width):
    if text is None:
        return ""
    text = str(text).replace("\n", " ").replace("\r", " ").strip()
    if len(text) <= width:
        return text
    return text[: width - 3] + "..."


def style_header(ws, row=1):
    """Bold + filled header row, frozen panes, autofilter."""
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="305496")
    for cell in ws[row]:
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="left", vertical="center")
    ws.freeze_panes = ws.cell(row=row + 1, column=1)
    ws.auto_filter.ref = ws.dimensions


def autosize_columns(ws, max_width=60):
    for col_idx, column_cells in enumerate(ws.columns, start=1):
        longest = 0
        for cell in column_cells:
            v = cell.value
            if v is None:
                continue
            longest = max(longest, len(str(v)))
        width = min(longest + 2, max_width)
        ws.column_dimensions[get_column_letter(col_idx)].width = width


def export_knowledge(conn, wb):
    ws = wb.create_sheet("Knowledge Base")
    ws.append(["ID", "Subject", "Answer (truncated)", "Timestamp"])

    rows = conn.execute(
        "SELECT id, subject, summary, timestamp FROM knowledge_base ORDER BY id"
    ).fetchall()

    for rid, subject, summary, ts in rows:
        ws.append([
            rid,
            truncate(subject, 200),
            truncate(summary, ANSWER_TRUNCATE),
            ts or "",
        ])

    style_header(ws)
    autosize_columns(ws, max_width=80)
    print(f"  Knowledge Base : {len(rows)} rows")
    return len(rows)


def export_moltbook(conn, wb):
    ws = wb.create_sheet("Moltbook Queries")
    ws.append(["ID", "Title / Question", "Post ID", "Verification Status",
               "DB Status", "Timestamp", "URL"])

    # Handle both schema variants (with/without verification_status)
    try:
        rows = conn.execute(
            "SELECT id, question, post_id, verification_status, status, timestamp "
            "FROM moltbook_queries ORDER BY id"
        ).fetchall()
    except sqlite3.OperationalError:
        rows = [
            (r[0], r[1], r[2], "", r[3], r[4])
            for r in conn.execute(
                "SELECT id, question, post_id, status, timestamp "
                "FROM moltbook_queries ORDER BY id"
            ).fetchall()
        ]

    for rid, title, post_id, vstatus, dstatus, ts in rows:
        url = f"https://www.moltbook.com/post/{post_id}" if post_id else ""
        ws.append([
            rid,
            truncate(title, TITLE_TRUNCATE),
            post_id or "",
            vstatus or "",
            dstatus or "",
            ts or "",
            url,
        ])

    style_header(ws)
    autosize_columns(ws, max_width=80)
    print(f"  Moltbook       : {len(rows)} rows")
    return len(rows)


def export_summary(wb, kb_count, mq_count, db_path):
    ws = wb.create_sheet("Summary", 0)
    ws.append(["Export Summary"])
    ws["A1"].font = Font(bold=True, size=14)
    ws.append([])
    ws.append(["Generated", datetime.now().strftime("%Y-%m-%d %H:%M:%S")])
    ws.append(["Database", db_path])
    ws.append(["Knowledge Base rows", kb_count])
    ws.append(["Moltbook queries", mq_count])
    ws.append(["Answer truncation (chars)", ANSWER_TRUNCATE])
    ws.append(["Title truncation (chars)", TITLE_TRUNCATE])
    ws.column_dimensions["A"].width = 32
    ws.column_dimensions["B"].width = 60


def main():
    if not os.path.exists(DB_NAME):
        print(f"❌ Database not found: {os.path.abspath(DB_NAME)}")
        print("   Run this from the folder that contains agent_memory.db")
        return

    db_path = os.path.abspath(DB_NAME)
    print(f"[Export] Reading {db_path}")

    conn = sqlite3.connect(DB_NAME)
    try:
        wb = Workbook()
        # Remove the default blank sheet
        wb.remove(wb.active)

        kb_count = export_knowledge(conn, wb)
        mq_count = export_moltbook(conn, wb)
        export_summary(wb, kb_count, mq_count, db_path)

        # Move Summary to the front
        wb.move_sheet("Summary", offset=-len(wb.sheetnames))
    finally:
        conn.close()

    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    out_name = f"agent_memory_export_{timestamp}.xlsx"
    out_path = os.path.join(os.path.dirname(db_path), out_name)

    wb.save(out_path)
    print(f"[Export] ✅ Saved: {out_path}")
    print(f"[Export]    Sheets: Summary | Knowledge Base | Moltbook Queries")


if __name__ == "__main__":
    main()